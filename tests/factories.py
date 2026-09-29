"""Generate real sample documents for extraction tests."""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


def make_pdf(path: Path, lines: list[str]) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=12)
        y += 20
    doc.save(str(path))
    doc.close()
    return path


def make_image_pdf(path: Path, text: str) -> Path:
    """A 'scanned' PDF: a single page containing only an image of text."""
    img_path = path.with_suffix(".png")
    make_png(img_path, text)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(str(path))
    doc.close()
    img_path.unlink()
    return path


def make_png(path: Path, text: str) -> Path:
    img = Image.new("RGB", (1400, 260), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=64)
    draw.text((40, 80), text, fill="black", font=font)
    img.save(path)
    return path


def make_docx(path: Path, paragraphs: list[str]) -> Path:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Week"
    table.rows[0].cells[1].text = "SQL joins"
    d.save(str(path))
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]]) -> Path:
    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        box.text_frame.text = body
    prs.save(str(path))
    return path
