from __future__ import annotations

from pathlib import Path

import docx
from pptx import Presentation

from app.processing.extractors.base import ExtractionError, ExtractionResult


def extract_docx(path: Path) -> ExtractionResult:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Cannot open DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ExtractionResult(text="\n".join(parts), method="docx")


def extract_pptx(path: Path) -> ExtractionResult:
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(f"Cannot open PPTX: {exc}") from exc
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        slide_parts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(notes)
        if slide_parts:
            parts.append(f"[{index}]\n" + "\n".join(slide_parts))
    return ExtractionResult(text="\n\n".join(parts), method="pptx", pages=len(prs.slides))
